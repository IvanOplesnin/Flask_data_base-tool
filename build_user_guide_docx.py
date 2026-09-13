"""Build the user guide DOCX from the reviewed Markdown source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "USER_GUIDE.md"
OUTPUT = ROOT / "USER_GUIDE.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_border(cell, color=GRID):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.add_run("Стр. ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def remove_paragraph_borders(paragraph):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn("w:pBdr"))
    if borders is not None:
        props.remove(borders)


def clean_text(text):
    text = text.strip()
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.replace("<http://217.26.31.67:8080/login>", "http://217.26.31.67:8080/login")


def add_inline(paragraph, text):
    """Minimal Markdown inline formatter: bold and code spans."""
    text = clean_text(text)
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def make_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_props = title._element.find(qn("w:pPr"))
    if title_props is not None:
        border = title_props.find(qn("w:pBdr"))
        if border is not None:
            title_props.remove(border)

    for level, size in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14 if level < 3 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_table(doc, rows):
    values = [[clean_text(cell) for cell in row] for row in rows]
    table = doc.add_table(rows=1, cols=len(values[0]))
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for col, value in enumerate(values[0]):
        cell = header.cells[col]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_no, values_row in enumerate(values[1:], start=1):
        cells = table.add_row().cells
        for col, value in enumerate(values_row):
            cell = cells[col]
            if row_no % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        raw = lines[index].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_markdown(doc, markdown):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            continue
        if stripped.startswith("#"):
            depth = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[depth:].strip()
            if depth == 1:
                # The cover title is created separately.
                i += 1
                continue
            style = f"Heading {min(depth - 1, 4)}"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, text)
            i += 1
            continue
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            add_inline(paragraph, stripped[1:])
            for run in paragraph.runs:
                run.bold = True
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet:
            # Avoid the inherited Word list styles: they can visually merge
            # neighbouring bullets when a document contains many sections.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run("• ")
            add_inline(paragraph, bullet.group(1))
            i += 1
            continue
        if numbered:
            # Word's built-in numbered-list style continues numbering between
            # unrelated Markdown lists. Keep the numbers from the source so
            # every separate instruction reliably starts with step 1.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(f"{numbered.group(1)}. ")
            add_inline(paragraph, numbered.group(2))
            i += 1
            continue
        if stripped.startswith("`") and stripped.endswith("`"):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(stripped.strip("`"))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt == "---" or nxt.startswith(("#", "|", ">", "- ", "* ")) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    make_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Инструкция по использованию приложения")
    remove_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Практическое руководство для пользователей").italic = True
    doc.add_paragraph()
    opening = doc.add_paragraph()
    opening.alignment = WD_ALIGN_PARAGRAPH.LEFT
    opening.add_run("Что вы найдёте в инструкции. ").bold = True
    opening.add_run(
        "Пошаговые действия для входа, поиска данных, добавления записей, загрузки экспериментов, "
        "работы с ролями и решения распространённых проблем."
    )
    doc.add_page_break()

    add_markdown(doc, SOURCE.read_text(encoding="utf-8"))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Инструкция по использованию приложения | ")
    add_page_number(footer)
    doc.core_properties.title = "Инструкция по использованию приложения"
    doc.core_properties.subject = "Практическое руководство для пользователей"
    doc.core_properties.author = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
